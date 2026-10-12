"""R12-F: 只读公开输出审计测试。

用 openpyxl 在临时目录生成隔离 xlsx 夹具，不写入仓库、不读取真实成果；
审计模块从 dev/audit-outputs.py 动态加载，避免向产品写个人机器路径。
"""
import hashlib
import importlib.util
import tempfile
import unittest
from pathlib import Path

import openpyxl

MODULE_PATH = Path(__file__).resolve().parents[1] / "dev" / "audit-outputs.py"
spec = importlib.util.spec_from_file_location("audit_outputs", MODULE_PATH)
audit_outputs = importlib.util.module_from_spec(spec)
spec.loader.exec_module(audit_outputs)

COLUMNS = ["序号", "项目编码", "项目名称", "项目特征描述", "单位", "工程量", "品牌", "报价单价", "金额", "备注"]
TITLE = "示例园区 水泵询价表"
CONTACT = "联系方式：010-12345678"
LONG_CODE = "00012345678901234567"
SHA = "a" * 64


def supplier_manifest(contact=CONTACT, long_code="SB-001"):
    return {
        "complete": True,
        "profile": "supplier",
        "sheets": [{
            "name": "水泵_青源", "kind": "inquiry", "visibility": "visible", "columns": 10,
            "title": {"cell": "A1", "value": TITLE, "merge": "A1:J1"},
            "merges": ["A3:E3", "F3:J3", "A4:E4", "F4:J4", "B5:J5"],
            "fields": [
                {"cell": "A3", "label": "项目名称", "value": "项目名称：示例园区"},
                {"cell": "F3", "label": "单位名称", "value": "单位名称：示例建设公司"},
                {"cell": "A4", "label": "材料经办人", "value": "材料经办人：测试员"},
                {"cell": "F4", "label": "联系方式", "value": contact, "contact": True},
            ],
            "supplierRegion": [
                {"cell": "A5", "value": "供应商信息"},
                {"cell": "B5", "value": "青源 — 华川设备"},
            ],
            "material": {
                "headerRow": 6, "firstDataRow": 7, "header": COLUMNS,
                "rows": [
                    ["1", long_code, "水泵", "DN50", "台", "2", "青源", "", "", ""],
                    ["2", "SB-002", "水泵", "DN65", "台", "3", "青源", "", "", ""],
                ],
            },
        }],
    }


def write_supplier(path, contact=CONTACT, long_code="SB-001", merges=True, hidden_column=None,
                   extra=None, internal=False, hidden_sheet=False, formula=None, text_equals=None):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "水泵_青源"
    sheet["A1"] = TITLE
    sheet["A3"] = "项目名称：示例园区"
    sheet["F3"] = "单位名称：示例建设公司"
    sheet["A4"] = "材料经办人：测试员"
    sheet["F4"] = contact
    sheet["A5"] = "供应商信息"
    sheet["B5"] = "青源 — 华川设备"
    if merges:
        for area in ("A1:J1", "A3:E3", "F3:J3", "A4:E4", "F4:J4", "B5:J5"):
            sheet.merge_cells(area)
    for index, header in enumerate(COLUMNS, start=1):
        sheet.cell(row=6, column=index, value=header)
    rows = [
        ["1", long_code, "水泵", "DN50", "台", "2", "青源", "", "", ""],
        ["2", "SB-002", "水泵", "DN65", "台", "3", "青源", "", "", ""],
    ]
    for row_index, row in enumerate(rows, start=7):
        for column_index, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column_index, value=value)
    if hidden_column:
        sheet.column_dimensions["K"].hidden = True
        for cell, value in hidden_column.items():
            sheet[cell] = value
    if extra:
        for cell, value in extra.items():
            sheet[cell] = value
    if formula:
        for cell, value in formula.items():
            sheet[cell] = value
    if text_equals:
        for cell, value in text_equals.items():
            sheet[cell] = value
            sheet[cell].data_type = "s"
    if internal:
        internal_sheet = workbook.create_sheet("处理记录")
        for row_index, row in enumerate([["类别", "品牌", "方式"], ["水泵", "青源", "合成一份"]], start=1):
            for column_index, value in enumerate(row, start=1):
                internal_sheet.cell(row=row_index, column=column_index, value=value)
    if hidden_sheet:
        hidden = workbook.create_sheet("临时隐藏页")
        hidden["A1"] = "x"
        hidden.sheet_state = "hidden"
    workbook.save(path)
    return path


def internal_manifest():
    manifest = supplier_manifest()
    manifest["profile"] = "internal-total"
    manifest["sheets"].append({
        "name": "处理记录", "kind": "internal", "visibility": "visible", "columns": 3,
        "material": {"headerRow": 1, "firstDataRow": 2, "header": ["类别", "品牌", "方式"],
                     "rows": [["水泵", "青源", "合成一份"]]},
    })
    return manifest


class AuditOutputsTest(unittest.TestCase):
    def setUp(self):
        self._temp = tempfile.TemporaryDirectory()
        self.addCleanup(self._temp.cleanup)
        self.dir = Path(self._temp.name)

    def test_project_code_and_joined_codes_do_not_become_phone_leaks(self):
        for index, code in enumerate(["030402011015", "030402011015；030402011017"]):
            path = write_supplier(self.dir / f"code-{index}.xlsx", long_code=code)
            report = audit_outputs.audit(str(path), supplier_manifest(long_code=code), profile="supplier")
            self.assertTrue(report["ok"], report["personal"])
            self.assertEqual(report["personal"], [])

    def test_code_exception_still_rejects_changed_values_headers_copies_and_phones(self):
        code = "030402011015"
        for scenario in ["changed-value", "wrong-header", "remark-copy", "phone-in-code"]:
            manifest = supplier_manifest(long_code=code)
            material = manifest["sheets"][0]["material"]
            actual, extra, cell = code, {}, "B7"
            if scenario == "changed-value": actual = "030402011017"
            if scenario == "wrong-header":
                extra["B6"] = "电话"
                material["header"] = list(material["header"])
                material["header"][1] = "电话"
            if scenario == "remark-copy":
                extra["J7"] = code
                material["rows"][0][9] = code
                cell = "J7"
            if scenario == "phone-in-code":
                actual = "010-12345678"
                material["rows"][0][1] = actual
            path = write_supplier(self.dir / (scenario+".xlsx"), long_code=actual, extra=extra)
            report = audit_outputs.audit(str(path), manifest, profile="supplier")
            self.assertFalse(report["ok"], scenario)
            self.assertTrue(any(x["cell"] == cell for x in report["personal"]), scenario)

    def test_supplier_file_with_allowed_contact_and_long_code_passes(self):
        path = write_supplier(self.dir / "ok.xlsx", long_code=LONG_CODE)
        report = audit_outputs.audit(str(path), supplier_manifest(long_code=LONG_CODE), profile="supplier")
        self.assertTrue(report["ok"], report)
        self.assertTrue(report["valueChecked"])
        self.assertEqual(report["verdict"], "value-verified")
        self.assertEqual(report["personal"], [])
        self.assertEqual(report["formulas"], [])
        self.assertEqual(report["sha256"], hashlib.sha256(path.read_bytes()).hexdigest())
        detail = report["sheetsDetail"][0]
        self.assertEqual(detail["name"], "水泵_青源")
        self.assertEqual(detail["rows"], 8)
        self.assertEqual(detail["columns"], 10)
        self.assertEqual(detail["formulaCount"], 0)

    def test_extra_hidden_column_with_phone_path_sha_and_price_is_rejected_by_cell(self):
        path = write_supplier(self.dir / "leak.xlsx", hidden_column={
            "K7": "010-12345678",
            "K8": r"Z:\机密\原清单.json",
            "K9": SHA,
            "K10": "原价 12.3 元",
        })
        report = audit_outputs.audit(str(path), supplier_manifest(), profile="supplier")
        self.assertFalse(report["ok"])
        self.assertEqual(report["hiddenColumns"], [{"sheet": "水泵_青源", "column": 11}])
        self.assertTrue(any(item["cell"] == "K7" for item in report["personal"]), report["personal"])
        self.assertTrue(any(item["cell"] == "K8" for item in report["paths"]), report["paths"])
        self.assertTrue(any(item["cell"] == "K9" for item in report["sha"]), report["sha"])
        self.assertTrue(any(item["cell"] == "K10" and item["field"] == "原价" for item in report["forbidden"]),
                        report["forbidden"])
        for cell in ("K7", "K8", "K9", "K10"):
            self.assertTrue(any(item["cell"] == cell for item in report["extraCells"]), cell)

    def test_extra_unexpected_value_and_missing_value_both_fail(self):
        extra = write_supplier(self.dir / "extra.xlsx", extra={"J7": "未预期备注"})
        report = audit_outputs.audit(str(extra), supplier_manifest(), profile="supplier")
        self.assertFalse(report["ok"])
        self.assertTrue(any(item["cell"] == "J7" for item in report["extraCells"]))

        missing = write_supplier(self.dir / "missing.xlsx", extra={"F7": ""})
        report = audit_outputs.audit(str(missing), supplier_manifest(), profile="supplier")
        self.assertFalse(report["ok"])
        hit = [item for item in report["missingCells"] if item["cell"] == "F7"]
        self.assertEqual(len(hit), 1)
        self.assertEqual(hit[0]["expected"], "2")

    def test_formula_internal_and_hidden_sheets_fail_and_text_equals_needs_range_scan(self):
        polluted = write_supplier(self.dir / "formula.xlsx", internal=True, hidden_sheet=True,
                                  formula={"H7": "=SUM(1+1)"})
        report = audit_outputs.audit(str(polluted), supplier_manifest(), profile="supplier")
        self.assertFalse(report["ok"])
        self.assertTrue(any(item["cell"] == "H7" for item in report["formulas"]))
        self.assertIn("处理记录", report["internalSheets"])
        self.assertIn("处理记录", report["internalSheetViolations"])
        self.assertIn("临时隐藏页", report["hiddenSheets"])

        text_file = write_supplier(self.dir / "text.xlsx", text_equals={"J7": "=1+1"})
        manifest = supplier_manifest()
        manifest["sheets"][0]["material"]["rows"][0][9] = {"value": "=1+1", "type": "text"}
        report = audit_outputs.audit(str(text_file), manifest, profile="supplier")
        self.assertEqual(report["formulas"], [], "文本=…不得误报为公式")
        self.assertTrue(report["ok"], report)

    def test_internal_total_allows_declared_internal_sheet_but_supplier_does_not(self):
        path = write_supplier(self.dir / "internal.xlsx", internal=True)
        allowed = audit_outputs.audit(str(path), internal_manifest(), profile="internal-total")
        self.assertTrue(allowed["ok"], allowed)
        self.assertEqual(allowed["internalSheetViolations"], [])
        self.assertIn("处理记录", allowed["internalSheets"])

        as_supplier = audit_outputs.audit(str(path), supplier_manifest(), profile="supplier")
        self.assertFalse(as_supplier["ok"])
        self.assertIn("处理记录", as_supplier["internalSheetViolations"])

        undeclared = audit_outputs.audit(str(path), supplier_manifest(), profile="internal-total")
        self.assertFalse(undeclared["ok"])
        self.assertIn("处理记录", undeclared["internalSheetViolations"])

    def test_scan_only_without_complete_manifest_never_claims_value_check(self):
        path = write_supplier(self.dir / "scan.xlsx", contact="联系方式：示例联系方式")
        report = audit_outputs.audit(str(path), None, profile="supplier")
        self.assertFalse(report["ok"])
        self.assertFalse(report["valueChecked"])
        self.assertEqual(report["verdict"], "scan-only")
        self.assertTrue(any("仅报告扫描结果" in note for note in report["notes"]), report["notes"])
        self.assertTrue(report["scanClean"])

        with_phone = audit_outputs.audit(str(write_supplier(self.dir / "scan-phone.xlsx")), None, profile="supplier")
        self.assertFalse(with_phone["valueChecked"])
        self.assertNotEqual(with_phone["verdict"], "value-verified")
        self.assertTrue(any(item["cell"] == "F4" for item in with_phone["personal"]),
                        "没有完整清单时不得豁免联系方式")

    def test_material_columns_and_layout_are_reported(self):
        path = write_supplier(self.dir / "layout.xlsx", extra={"K6": "额外列"})
        report = audit_outputs.audit(str(path), supplier_manifest(), profile="supplier")
        self.assertFalse(report["ok"])
        self.assertTrue(any(item["cell"] == "K6" for item in report["extraCells"]))
        self.assertTrue(any(item["code"] == "column-beyond-declared" for item in report["layoutIssues"]))

    def test_material_rows_follow_declared_header_position_not_rise_header(self):
        path = write_supplier(self.dir / "header.xlsx")
        manifest = supplier_manifest()
        manifest["sheets"][0]["material"]["headerRow"] = 1
        manifest["sheets"][0]["material"]["firstDataRow"] = 2
        report = audit_outputs.audit(str(path), manifest, profile="supplier")
        self.assertFalse(report["ok"])
        self.assertTrue(any(item["cell"] == "A1" for item in report["mismatches"]), report["mismatches"])
        self.assertTrue(any(item["cell"] == "A2" for item in report["missingCells"]), report["missingCells"])


if __name__ == "__main__":
    unittest.main()
