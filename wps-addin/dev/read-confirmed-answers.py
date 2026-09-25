"""只读读取人工核对表，输出可供 buildMigration 使用的 rows/overrides。

用法：python read-confirmed-answers.py <核对表xlsx> <answer-overrides.json> <输出json>
只读输入文件，不修改任何原始资料；核对人姓名、电话、邮箱等个人数据不进输出。
"""
import hashlib
import json
import re
import sys

import openpyxl

SHEET = "核对填写"
HEADER_ROW = 4
COLUMNS = {"name": 2, "feature": 3, "unit": 4, "decision": 5, "category": 6,
           "reviewer": 7, "date": 8, "basis": 9, "project": 1, "record": 22}
PERSONAL = re.compile(r"(?:\d[\d\-\s]{6,}\d|[\w.+-]+@[\w-]+\.[\w.]+)")
DECISION_SKIP = "无需询价"
SKIP_CATEGORY = "无需询价"


def text(value):
    return "" if value is None else str(value).strip()


def safe(text_value):
    """确认依据只保留业务口径文字，去掉电话、邮箱等个人联系方式。"""
    return PERSONAL.sub("[已隐去]", text_value)[:60]


def sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_rows(workbook):
    sheet = workbook[SHEET]
    rows, rejected = [], {"unconfirmed": 0, "no-category": 0, "no-name": 0}
    for index, values in enumerate(sheet.iter_rows(min_row=HEADER_ROW + 1, max_col=25, values_only=True), start=HEADER_ROW + 1):
        cells = [text(value) for value in values]
        if not any(cells):
            continue
        record_id = cells[COLUMNS["record"]] or ("row-%d" % index)
        decision, category = cells[COLUMNS["decision"]], cells[COLUMNS["category"]]
        confirmed = all(cells[COLUMNS[key]] for key in ("decision", "reviewer", "date", "basis"))
        if not confirmed:
            rejected["unconfirmed"] += 1
            continue
        name = cells[COLUMNS["name"]]
        if not name:
            rejected["no-name"] += 1
            continue
        if not category and decision == DECISION_SKIP:
            category = SKIP_CATEGORY
        if not category:
            rejected["no-category"] += 1
            continue
        rows.append({
            "recordId": record_id,
            "project": cells[COLUMNS["project"]],
            "excelRow": index,
            "item": {"name": name, "feature": cells[COLUMNS["feature"]], "unit": cells[COLUMNS["unit"]]},
            "categoryId": None,
            "categoryName": category,
            "confirmed": True,
            "confirmationSource": safe(cells[COLUMNS["basis"]]) or "人工核对表",
        })
    return rows, rejected


def main():
    source, overrides_path, output = sys.argv[1], sys.argv[2], sys.argv[3]
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    rows, rejected = read_rows(workbook)
    extra = workbook["附加资料"]
    extra_rows = sum(1 for values in extra.iter_rows(min_row=HEADER_ROW + 1, max_col=8, values_only=True)
                     if any(text(value) for value in values))
    with open(overrides_path, encoding="utf-8") as handle:
        overrides = json.load(handle)
    payload = {"source": {"path": source, "sha256": sha256(source), "sheet": SHEET, "headerRow": HEADER_ROW},
               "counts": {"rows": len(rows), "extraRows": extra_rows, "overrides": len(overrides)}, "rejected": rejected,
               "rows": rows, "overrides": overrides}
    with open(output, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=1)
    print(json.dumps({"rows": len(rows), "extraRows": extra_rows, "overrides": len(overrides),
                      "rejected": rejected, "output": output}, ensure_ascii=False))


if __name__ == "__main__":
    main()
