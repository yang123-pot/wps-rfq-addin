"""R12-F 公开输出审计：只读读取实际 xlsx，按完整预期清单逐值/版式/额外内容核对。

用法：
    python audit-outputs.py [--profile supplier|internal-total] [--manifest manifest.json] <xlsx...>

未提供 --manifest（或清单 complete!=true）时只报告扫描结果（valueChecked=false，
verdict='scan-only'），不会给出“逐值通过”结论。审计不修改任何文件。

预期清单格式（与 dev/release-audit.cjs 的 expectedManifest 一致）：
{
  "complete": true,
  "profile": "supplier" | "internal-total",
  "sheets": [{
    "name": "水泵_青源", "kind": "inquiry" | "internal", "visibility": "visible",
    "columns": 10,
    "title": {"cell": "A1", "value": "...", "merge": "A1:J1"},
    "fields": [{"cell": "F4", "label": "联系方式", "value": "联系方式：010-12345678", "contact": true}],
    "supplierRegion": [{"cell": "A5", "value": "供应商信息"}],
    "merges": ["A3:E3", "F3:J3"],
    "material": {"headerRow": 6, "firstDataRow": 7, "header": ["序号", "..."],
                 "rows": [["1", "SB-001", "..."]]},
    "cells": [{"cell": "J7", "value": "备注"}]
  }]
}
"""
import argparse
import hashlib
import json
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string

PROFILES = ["supplier", "internal-total"]
FORBIDDEN = ["原价", "历史报价", "原清单文件", "原路径", "SHA256", "核对追溯JSON", "自动结果", "历史分类定位"]
INTERNAL = ["内部追溯", "处理记录", "来源明细", "清单"]

CONTACT_RE = re.compile(r"(?:^|[^0-9])(?:1[0-9]{10}|0[0-9]{2,3}-?[0-9]{7,8}(?:-[0-9]{1,5})?)(?=$|[^0-9])")
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PATH_RE = re.compile(r"(?:^|[^A-Za-z0-9_])(?:[A-Za-z]:[\\/]|\\\\[^\\\s]+[\\/]|file://)")
SHA_RE = re.compile(r"(?:^|[^0-9A-Za-z])[0-9a-fA-F]{40,64}(?=$|[^0-9A-Za-z])")
ADDRESS_RE = re.compile(r"^([A-Za-z]+)([0-9]+)$")


def text(value):
    return "" if value is None else str(value)


def _sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_address(address):
    match = ADDRESS_RE.match(text(address).strip())
    if not match:
        return None
    letters = match.group(1).upper()
    column = 0
    for char in letters:
        column = column * 26 + (ord(char) - 64)
    return {"row": int(match.group(2)), "column": column}


def normalize_range(value):
    if isinstance(value, dict):
        value = value.get("range") or value.get("cell") or ""
    return text(value).upper().replace("$", "")


def range_top_left(value):
    return normalize_range(value).split(":")[0]


def normalize_type(data_type, is_formula):
    if is_formula:
        return "formula"
    if data_type in ("s", "str", "inlineStr"):
        return "text"
    if data_type == "n":
        return "number"
    if data_type == "b":
        return "boolean"
    return data_type or "text"


def entry_value(raw):
    if isinstance(raw, dict) and ("value" in raw or "text" in raw):
        return {"value": raw.get("value", raw.get("text")), "type": raw.get("type")}
    return {"value": raw, "type": None}


def expand_expected(sheet_manifest):
    """把清单的区域声明展开成 {单元格: 预期条目}。"""
    cells = {}

    def add(cell, value, **options):
        position = parse_address(cell)
        if not position or value is None or value == "":
            return
        cells[cell.upper()] = dict(
            cell=cell.upper(), row=position["row"], column=position["column"], value=text(value), **options)

    title = sheet_manifest.get("title")
    if title and title.get("value") is not None:
        add(title.get("cell") or range_top_left(title.get("range")), title.get("value"), role="title")
    for field in sheet_manifest.get("fields", []):
        add(field.get("cell"), field.get("value"), role="field",
            contact=field.get("contact") is True, type=field.get("type"))
    for field in sheet_manifest.get("supplierRegion", []):
        add(field.get("cell"), field.get("value"), role="supplier",
            contact=field.get("contact") is True, type=field.get("type"))
    for field in sheet_manifest.get("cells", []):
        add(field.get("cell"), field.get("value"), role=field.get("role"),
            contact=field.get("contact") is True, type=field.get("type"))
    material = sheet_manifest.get("material")
    if material:
        header_row = int(material.get("headerRow"))
        first_data_row = int(material.get("firstDataRow"))
        for index, raw in enumerate(material.get("header", [])):
            item = entry_value(raw)
            add(_address(header_row, index + 1), item["value"], role="material-header", type=item["type"])
        for row_index, row in enumerate(material.get("rows", [])):
            for column_index, raw in enumerate(row):
                item = entry_value(raw)
                headers = material.get("header", [])
                code_column = column_index < len(headers) and text(entry_value(headers[column_index])["value"]) == "项目编码"
                add(_address(first_data_row + row_index, column_index + 1), item["value"],
                    role="material", type=item["type"], projectCode=code_column)
    return cells


def _address(row, column):
    letters = ""
    rest = int(column)
    while rest > 0:
        remainder = (rest - 1) % 26
        letters = chr(65 + remainder) + letters
        rest = (rest - 1) // 26
    return "%s%d" % (letters, int(row))


def expected_merges(sheet_manifest):
    ranges = []
    title = sheet_manifest.get("title") or {}
    if title.get("merge"):
        ranges.append(normalize_range(title["merge"]))
    for field in sheet_manifest.get("fields", []):
        if field.get("merge"):
            ranges.append(normalize_range(field["merge"]))
    for field in sheet_manifest.get("supplierRegion", []):
        if field.get("merge"):
            ranges.append(normalize_range(field["merge"]))
    for raw in sheet_manifest.get("merges", []):
        ranges.append(normalize_range(raw))
    return [item for item in ranges if item]


def hidden_columns(sheet):
    columns = set()
    for key, dimension in sheet.column_dimensions.items():
        if not getattr(dimension, "hidden", False):
            continue
        if isinstance(key, str) and key.isalpha():
            columns.add(column_index_from_string(key))
        else:
            start = dimension.min or 1
            end = dimension.max or start
            for index in range(start, end + 1):
                columns.add(index)
    return sorted(columns)


def collect_cells(sheet):
    cells = {}
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if value is None or value == "":
                continue
            is_formula = cell.data_type == "f"
            cells[cell.coordinate] = {
                "value": value, "text": text(value), "isFormula": is_formula,
                "dataType": normalize_type(cell.data_type, is_formula),
            }
    return cells


def _pattern_scan(sheet_name, cell_ref, info, expected, report):
    position = parse_address(cell_ref)
    if info["isFormula"]:
        report["formulas"].append({"sheet": sheet_name, "cell": cell_ref,
                                   "row": position["row"], "column": position["column"], "text": info["text"]})
    for word in FORBIDDEN:
        if word in info["text"]:
            report["forbidden"].append({"sheet": sheet_name, "cell": cell_ref, "row": position["row"],
                                        "column": position["column"], "field": word})
    if PATH_RE.search(info["text"]):
        report["paths"].append({"sheet": sheet_name, "cell": cell_ref, "row": position["row"],
                                "column": position["column"], "text": info["text"]})
    if SHA_RE.search(info["text"]):
        report["sha"].append({"sheet": sheet_name, "cell": cell_ref, "row": position["row"],
                              "column": position["column"], "text": info["text"]})
    if CONTACT_RE.search(info["text"]) or EMAIL_RE.search(info["text"]):
        # 编码字段语义＋预期原值＋12位编码格式三者齐备，才排除与座机正则重合的编码。
        allowed_code = bool(expected and expected.get("projectCode") is True and expected.get("value") == info["text"]
                            and re.fullmatch(r"[0-9]{12}(?:；[0-9]{12})*", info["text"]))
        allowed = allowed_code or bool(expected and expected.get("contact") is True and expected.get("value") == info["text"])
        if not allowed:
            report["personal"].append({"sheet": sheet_name, "cell": cell_ref, "row": position["row"],
                                       "column": position["column"], "text": info["text"]})


def audit(path, expected_manifest=None, profile=None):
    """只读审计一个 xlsx。expected_manifest 完整时才给出逐值结论。"""
    manifest = expected_manifest or None
    manifest_profile = manifest.get("profile") if isinstance(manifest, dict) else None
    resolved_profile = "internal-total" if text(profile or manifest_profile) == "internal-total" else "supplier"
    manifest_complete = bool(isinstance(manifest, dict) and manifest.get("complete") is True
                             and isinstance(manifest.get("sheets"), list))
    expected_sheets = {}
    if manifest_complete:
        for item in manifest.get("sheets", []):
            expected_sheets[text(item.get("name"))] = item

    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    report = {
        "path": path,
        "sha256": _sha256(path),
        "profile": resolved_profile,
        "ok": False,
        "valueChecked": False,
        "manifestComplete": manifest_complete,
        "verdict": "failed",
        "scanClean": False,
        "notes": [],
        "sheets": [],
        "sheetsDetail": [],
        "scan": {"sheets": 0, "formulaCount": 0, "commentCount": 0, "hyperlinkCount": 0},
        "hiddenSheets": [],
        "internalSheets": [],
        "internalSheetViolations": [],
        "unexpectedSheets": [],
        "missingSheets": [],
        "hiddenColumns": [],
        "formulas": [],
        "forbidden": [],
        "personal": [],
        "paths": [],
        "sha": [],
        "commentCells": [],
        "hyperlinkCells": [],
        "externalLinks": [],
        "externalRelationships": len(getattr(workbook, "_external_links", []) or []),
        "extraCells": [],
        "missingCells": [],
        "mismatches": [],
        "typeIssues": [],
        "layoutIssues": [],
    }
    actual_names = [sheet.title for sheet in workbook.worksheets]
    report["sheets"] = actual_names

    for sheet in workbook.worksheets:
        name = sheet.title
        cells = collect_cells(sheet)
        merges = [str(item).upper() for item in sheet.merged_cells.ranges]
        expected_sheet = expected_sheets.get(name)
        expected_cells = expand_expected(expected_sheet) if expected_sheet else None
        declared_internal = resolved_profile == "internal-total" and bool(expected_sheet) and expected_sheet.get("kind") == "internal"

        if sheet.sheet_state != "visible":
            report["hiddenSheets"].append(name)
        internal = sheet.title in INTERNAL or (expected_sheet or {}).get("kind") == "internal"
        if internal:
            report["internalSheets"].append(name)
            if not declared_internal:
                report["internalSheetViolations"].append(name)
        if manifest_complete and not expected_sheet:
            report["unexpectedSheets"].append({"sheet": name, "code": "unexpected-sheet", "message": "清单之外的工作表：" + name})

        declared_hidden = []
        if expected_sheet:
            for raw in expected_sheet.get("hiddenColumns", []):
                parsed = parse_address(text(raw) + "1") if isinstance(raw, str) else None
                declared_hidden.append(parsed["column"] if parsed else int(raw))
        for column in hidden_columns(sheet):
            if not (resolved_profile == "internal-total" and column in declared_hidden):
                report["hiddenColumns"].append({"sheet": name, "column": column})

        declared_columns = int(expected_sheet.get("columns") or 0) if expected_sheet else 0
        for cell_ref, info in cells.items():
            position = parse_address(cell_ref)
            expected = expected_cells.get(cell_ref) if expected_cells else None
            if not declared_internal:
                _pattern_scan(name, cell_ref, info, expected, report)
            if expected_cells is not None:
                if expected is None:
                    report["extraCells"].append({"sheet": name, "cell": cell_ref, "row": position["row"],
                                                 "column": position["column"], "actual": info["text"]})
                else:
                    if expected["value"] != info["text"]:
                        report["mismatches"].append({"sheet": name, "cell": cell_ref, "row": position["row"],
                                                     "column": position["column"], "expected": expected["value"],
                                                     "actual": info["text"]})
                    if expected.get("type") == "text" and info["dataType"] not in ("text", "string"):
                        report["typeIssues"].append({"sheet": name, "cell": cell_ref, "row": position["row"],
                                                     "column": position["column"], "expected": expected["value"],
                                                     "actualType": info["dataType"]})
                if declared_columns and position["column"] > declared_columns:
                    report["layoutIssues"].append({"sheet": name, "code": "column-beyond-declared",
                                                   "detail": {"cell": cell_ref, "column": position["column"],
                                                              "declared": declared_columns}})

        if expected_cells is not None:
            for cell_ref, expected in expected_cells.items():
                if cell_ref not in cells:
                    report["missingCells"].append({"sheet": name, "cell": cell_ref, "row": expected["row"],
                                                   "column": expected["column"], "expected": expected["value"]})

        if expected_sheet and manifest_complete:
            declared_merges = expected_merges(expected_sheet)
            for item in declared_merges:
                if item not in merges:
                    report["layoutIssues"].append({"sheet": name, "code": "missing-merge", "detail": item})
            for item in merges:
                if item not in declared_merges:
                    report["layoutIssues"].append({"sheet": name, "code": "unexpected-merge", "detail": item})

        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    report["commentCells"].append({"sheet": name, "cell": cell.coordinate})
                if cell.hyperlink is not None:
                    target = text(getattr(cell.hyperlink, "target", "") or getattr(cell.hyperlink, "location", ""))
                    item = {"sheet": name, "cell": cell.coordinate, "target": target}
                    report["hyperlinkCells"].append(item)
                    if PATH_RE.search(target) or re.match(r"^https?://", target, re.I):
                        report["externalLinks"].append(item)

        rows = max([sheet.max_row] + [parse_address(key)["row"] for key in cells] + [1])
        columns = max([sheet.max_column] + [parse_address(key)["column"] for key in cells] + [1])
        for item in merges:
            corners = item.split(":")
            for corner in corners:
                position = parse_address(corner)
                if position:
                    rows = max(rows, position["row"])
                    columns = max(columns, position["column"])
        report["sheetsDetail"].append({
            "name": name, "rows": rows, "columns": columns,
            "formulaCount": sum(1 for info in cells.values() if info["isFormula"]),
            "state": "visible" if sheet.sheet_state == "visible" else "hidden",
        })

    if manifest_complete:
        for item in manifest.get("sheets", []):
            name = text(item.get("name"))
            if name not in actual_names:
                report["missingSheets"].append({"sheet": name, "code": "missing-sheet", "message": "文件里没有清单声明的表：" + name})

    report["scan"] = {
        "sheets": len(actual_names),
        "formulaCount": len(report["formulas"]),
        "commentCount": len(report["commentCells"]),
        "hyperlinkCount": len(report["hyperlinkCells"]),
    }
    violation_lists = ["hiddenSheets", "internalSheetViolations", "unexpectedSheets", "missingSheets", "hiddenColumns",
                       "formulas", "forbidden", "personal", "paths", "sha", "commentCells", "hyperlinkCells",
                       "externalLinks", "extraCells", "missingCells", "mismatches", "typeIssues", "layoutIssues"]
    violations = sum(len(report[key]) for key in violation_lists) + report["externalRelationships"]
    report["scanClean"] = violations == 0
    report["valueChecked"] = bool(manifest_complete and not report["missingSheets"] and not report["unexpectedSheets"])
    report["ok"] = bool(report["valueChecked"] and report["scanClean"])
    if report["ok"]:
        report["verdict"] = "value-verified"
    elif report["scanClean"] and not report["valueChecked"]:
        report["verdict"] = "scan-only"
    else:
        report["verdict"] = "failed"
    if not manifest_complete:
        report["notes"].append("未提供完整预期清单（manifest.complete!=true），仅报告扫描结果，不能判定逐值通过。")
    elif not report["valueChecked"]:
        report["notes"].append("预期清单未覆盖实际工作表或声明的表缺失，不能判定逐值通过。")
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description="只读审计公开输出的 xlsx")
    parser.add_argument("--profile", choices=PROFILES, default=None)
    parser.add_argument("--manifest", default=None, help="完整预期清单 JSON 路径")
    parser.add_argument("files", nargs="+")
    args = parser.parse_args(argv)
    manifest = None
    if args.manifest:
        with open(args.manifest, "r", encoding="utf-8") as handle:
            manifest = json.load(handle)
    reports = [audit(path, manifest, profile=args.profile) for path in args.files]
    print(json.dumps({"ok": all(item["ok"] for item in reports),
                      "valueChecked": all(item["valueChecked"] for item in reports),
                      "files": reports}, ensure_ascii=False, indent=1))
    return 0 if all(item["ok"] for item in reports) else 1


if __name__ == "__main__":
    sys.exit(main())
