"""Create an eight-page, primarily graphical WPS user guide.

Original UI screenshots are composed into instructional figures with arrows.
Editable headings and the EXE download link remain native DOCX.
"""
from pathlib import Path
import json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
DOCS=ROOT.parents[1]
SHOTS=DOCS/'images'/'wps-user-guide-v010-20261020'
FIGS=ROOT/'figures'
OUT=ROOT/'材料整理助手操作图解0.1.0.docx'
FONT='C:/Windows/Fonts/msyh.ttc'
BOLD='C:/Windows/Fonts/msyhbd.ttc'
GREEN='#22744f'; INK='#14251c'; GRAY='#59685f'; LINE='#ccd8d0'; PALE='#eff6f1'; YELLOW='#fff2b7'

def font(n=42,bold=False):return ImageFont.truetype(BOLD if bold else FONT,n)
def txt(d,x,y,s,n=42,bold=False,color=INK):d.multiline_text((x,y),s,font=font(n,bold),fill=color,spacing=17)
def mid(d,b,s,n=42,bold=False,color=INK):
    bb=d.multiline_textbbox((0,0),s,font=font(n,bold),spacing=16)
    d.multiline_text(((b[0]+b[2]-bb[2])/2,(b[1]+b[3]-(bb[3]-bb[1]))/2-bb[1]),s,font=font(n,bold),fill=color,spacing=16,align='center')
def box(d,b,s,n=42,fill='white',outline=GREEN,bold=False,color=INK):
    d.rounded_rectangle(b,20,fill=fill,outline=outline,width=4);mid(d,b,s,n,bold,color)
def arrow(d,a,b):
    import math
    d.line((a,b),fill=GREEN,width=8);q=math.atan2(b[1]-a[1],b[0]-a[0]);s=25;h=15
    d.polygon([b,(b[0]-s*math.cos(q)+h*math.sin(q),b[1]-s*math.sin(q)-h*math.cos(q)),(b[0]-s*math.cos(q)-h*math.sin(q),b[1]-s*math.sin(q)+h*math.cos(q))],fill=GREEN)
def number(d,x,y,n):
    d.ellipse((x,y,x+68,y+68),fill=GREEN);mid(d,(x,y-3,x+68,y+68),str(n),43,True,'white')
def step(d,x,y,n,label,size=46):number(d,x,y,n);txt(d,x+93,y+5,label,size,True)
def image(im,path,x,y,w,crop=None):
    pic=Image.open(path).convert('RGB')
    if crop:pic=pic.crop(crop)
    h=round(w*pic.height/pic.width);pic=pic.resize((w,h),Image.Resampling.LANCZOS);im.paste(pic,(x,y));return (x,y,x+w,y+h)
def shot(im,name,x,y,w,crop=None):return image(im,SHOTS/name,x,y,w,crop)
def button(im,name,x,y,w,crop=(85,714,386,760)):
    labels={'step-00-simple-home.png':'选择清单文件','step-01-import.png':'确认选择，开始自动分类','step-02-classification.png':'生成分类核对表','step-03-review.png':'核对完成，生成分类工作簿','step-05a-summary-separated.png':'下一步：生成询价表','step-07-split.png':'开始拆分','step-07-complete.png':'打开文件夹'}
    h=145
    box(ImageDraw.Draw(im),(x,y,x+w,y+h),labels[name],47,fill=GREEN,color='white',bold=True)
    return (x,y,x+w,y+h)
def sheet(d,x,y,w=160,h=180):
    d.rounded_rectangle((x,y,x+w,y+h),12,fill='white',outline=GREEN,width=5)
    d.rectangle((x+22,y+38,x+w-22,y+h-25),outline=GREEN,width=4)
    for k in (1,2):d.line((x+22,y+38+k*(h-63)/3,x+w-22,y+38+k*(h-63)/3),fill=GREEN,width=3)
    d.line((x+w/2,y+38,x+w/2,y+h-25),fill=GREEN,width=3)
def folder(d,x,y,w=200,h=150):
    d.rounded_rectangle((x,y,x+w*.45,y+45),10,fill='#dbeedf',outline=GREEN,width=4)
    d.rounded_rectangle((x,y+25,x+w,y+h),10,fill='#f3f8f4',outline=GREEN,width=5)
def new(h=2050):return Image.new('RGB',(1600,h),'white')
def save(i,im):
    p=FIGS/f'{i:02d}.png';im.save(p,optimize=True);return p

def page1():
    im=new();d=ImageDraw.Draw(im)
    txt(d,40,10,'已经装好  →  直接看下一页',43,True)
    step(d,40,130,1,'下载 EXE 安装包')
    sheet(d,100,250,140,155)
    box(d,(330,245,1510,390),'WpsRfqAddin-Setup-0.1.0.exe',39)
    txt(d,330,420,'约145MB  无需解压  无需输入命令',35,color=GRAY)
    step(d,40,560,2,'保存并关闭 WPS → 双击 EXE')
    step(d,40,730,3,'等待检测通过 → 点 一键安装')
    image(im,ROOT/'installer-interface.png',120,850,1030,crop=(0,0,670,330))
    txt(d,1200,1010,'运行环境\n已内置',38,True,color=GREEN)
    txt(d,1200,1160,'不用另外\n下载安装',35,color=GRAY)
    box(d,(430,1405,1230,1535),'一键安装',50,fill=GREEN,color='white',bold=True)
    txt(d,425,1570,'已装过时，按钮显示“升级 / 修复”',32,color=GRAY)
    step(d,40,1690,4,'安装完成 → 点 打开 WPS')
    box(d,(430,1795,1230,1920),'打开 WPS',48,fill=GREEN,color='white',bold=True)
    txt(d,120,1970,'缺少 WPS？点“安装 WPS”，装好后点“重新检测”。',31,color=GRAY)
    return im

def page2():
    im=new();d=ImageDraw.Draw(im)
    labels=['WPS表格','材料整理','打开材料整理助手']
    for i,label in enumerate(labels):
        x=30+i*535;box(d,(x,20,x+460,145),label,38,fill=PALE,bold=True)
        if i:arrow(d,(x-65,82),(x-10,82))
    b=shot(im,'step-00-simple-home.png',90,250,650)
    txt(d,95,185,'日常使用：也可双击桌面“材料整理助手”打开 WPS',35,color=GRAY)
    step(d,850,355,1,'看当前步骤',43);arrow(d,(845,385),(155,385))
    step(d,850,750,2,'按绿色按钮',43)
    txt(d,850,870,'照提示往下做',43)
    arrow(d,(1040,1100),(720,1440))
    txt(d,160,1635,'要继续以前的工作？',47,True)
    box(d,(160,1740,1440,1910),'材料整理 → 打开项目 → 选原项目',44,fill=PALE)
    return im

def page3():
    im=new();d=ImageDraw.Draw(im)
    step(d,40,20,1,'点 选择清单文件')
    button(im,'step-00-simple-home.png',270,125,1050,crop=(15,714,386,760))
    step(d,40,320,2,'选本项目的 .xls 或 .xlsx')
    for i,label in enumerate(['给排水清单.xlsx','机电清单.xlsx']):
        y=440+i*145;sheet(d,140,y,85,100);box(d,(290,y,1450,y+105),label,45)
        d.rectangle((330,y+31,376,y+77),outline=GREEN,width=4)
        d.line([(338,y+53),(348,y+65),(369,y+40)],fill=GREEN,width=5)
    txt(d,290,740,'多份文件：按住 Ctrl 再逐个点击',38,color=GRAY)
    step(d,40,855,3,'看 导入预览 → 核对文件和工作表')
    box(d,(150,980,1450,1110),'文件齐全、数量正确 → 我已看完预览',43,fill=PALE)
    step(d,40,1240,4,'选择是否 合并清单')
    box(d,(140,1360,1450,1470),'不勾选 → 按原清单逐条保留',43)
    box(d,(140,1510,1450,1660),'勾选 → 名称 特征 单位相同的材料\n数量相加',41,fill=YELLOW,outline='#bca348')
    arrow(d,(800,1710),(800,1790))
    button(im,'step-01-import.png',230,1820,1170,crop=(75,714,386,760))
    return im

def page4():
    im=new();d=ImageDraw.Draw(im)
    step(d,40,20,1,'生成并打开分类核对表')
    button(im,'step-02-classification.png',260,125,1080,crop=(15,714,386,760))
    step(d,40,360,2,'在黄色 最终分类 列里改分类')
    rows=[['项目名称','项目特征','最终分类'],['水泵','DN50','水泵'],['配电箱','充电设备配套','汽车充电桩'],['钢索','规格C','无需询价']]
    xs=[120,520,1050,1490]
    for ri,row in enumerate(rows):
        for ci,value in enumerate(row):
            y=510+ri*142
            d.rectangle((xs[ci],y,xs[ci+1],y+142),fill=PALE if ri==0 else YELLOW if ci==2 else 'white',outline=LINE,width=3)
            mid(d,(xs[ci],y,xs[ci+1],y+142),value,40,ri==0)
    arrow(d,(1390,1200),(1290,1030))
    txt(d,120,1210,'每条都填完整',47,True)
    txt(d,120,1315,'不询价的材料 → 填“无需询价”',44)
    txt(d,120,1410,'橙色提示 → 再核对一下分类',42,color=GRAY)
    step(d,40,1570,3,'核对完 → 点绿色按钮')
    button(im,'step-03-review.png',180,1700,1240,crop=(170,714,386,760))
    txt(d,260,1990,'按提示保存并生成分类工作簿',38,color=GRAY)
    return im

def page5():
    im=new();d=ImageDraw.Draw(im)
    step(d,40,20,1,'打开 分类工作簿 的第一个标签')
    tabs=[('分类汇总表',True),('水泵',False),('阀门',False),('无需询价',False)]
    x=120
    for value,selected in tabs:
        w=450 if selected else 300
        box(d,(x,160,x+w,265),value,40,fill=PALE if selected else 'white',bold=selected);x+=w
    arrow(d,(330,330),(330,275))
    step(d,40,400,2,'品牌 供应商 填在黄色区域')
    rows=[['名称','品牌一','供应商一'],['水泵','示例品牌甲','示例设备甲'],['阀门','','']]
    xs=[120,520,1000,1480]
    for ri,row in enumerate(rows):
        for ci,value in enumerate(row):
            y=555+ri*135
            d.rectangle((xs[ci],y,xs[ci+1],y+135),fill=PALE if ri==0 else YELLOW if ci>0 else 'white',outline=LINE,width=3)
            mid(d,(xs[ci],y,xs[ci+1],y+135),value,40,ri==0)
    txt(d,160,1020,'可以全部留空',57,True,color=GREEN)
    txt(d,160,1125,'有多组品牌时，再选生成方式',43)
    box(d,(120,1270,760,1450),'合成一份\n两组 → 一份询价',42)
    box(d,(840,1270,1480,1450),'按品牌分别生成\n两组 → 两份询价',40)
    txt(d,220,1505,'每份都是完整材料，工程量不按组数分摊',35,color=GRAY)
    step(d,40,1650,3,'Ctrl+S 保存 → 回助手点下一步')
    button(im,'step-05a-summary-separated.png',220,1800,1160,crop=(15,714,386,760))
    return im

def page6():
    im=new();d=ImageDraw.Draw(im)
    step(d,40,20,1,'填四项必填信息')
    b=shot(im,'step-06-inquiry.png',70,150,680)
    for label,source_y in [('项目名称',346),('单位名称',398),('经办人信息',450),('联系方式',502)]:
        center_y=round(150+source_y*1.7)
        txt(d,940,center_y-30,label,46,True);arrow(d,(915,center_y),(710,center_y))
    txt(d,925,1250,'其余三项\n可以留空',47,color=GRAY)
    step(d,40,1540,2,'点 生成询价总表')
    arrow(d,(800,1640),(800,1720))
    box(d,(130,1760,1470,1950),'选保存位置 → 输入新文件名 → 保存',44,fill=PALE)
    return im

def page7():
    im=new();d=ImageDraw.Draw(im)
    step(d,40,20,1,'选择一个空输出文件夹')
    folder(d,160,155,220,160)
    box(d,(490,160,1470,330),'例如  供应商文件_20261020',44)
    step(d,40,460,2,'点 开始拆分 → 等待完成')
    button(im,'step-07-split.png',200,575,1200)
    arrow(d,(800,820),(800,900))
    shot(im,'step-07-complete.png',220,930,1160,crop=(12,180,385,273))
    step(d,40,1310,3,'点 打开文件夹')
    button(im,'step-07-complete.png',200,1430,1200,crop=(85,448,386,493))
    sheet(d,140,1695,110,140)
    txt(d,320,1715,'核对后，把对应询价文件发给供应商',43,True)
    txt(d,320,1810,'“完整总表”和“无需询价”先留作内部记录',36,color=GRAY)
    txt(d,320,1920,'文件份数以完成页实际显示为准',36,color=GRAY)
    return im

def page8():
    im=new();d=ImageDraw.Draw(im)
    step(d,40,20,1,'暂停或关机前')
    box(d,(120,160,1480,290),'在WPS中 Ctrl+S → 保存表格',48,fill=PALE)
    arrow(d,(800,330),(800,395))
    box(d,(120,435,1480,580),'助手 设置 → 项目状态 → 保存项目',45)
    step(d,40,720,2,'下次继续')
    box(d,(120,850,1480,990),'双击桌面“材料整理助手”',45,fill=PALE)
    arrow(d,(800,1030),(800,1100))
    box(d,(120,1140,1480,1290),'WPS表格 → 材料整理 → 打开项目',44)
    txt(d,130,1415,'拆分成功后：',41,True)
    box(d,(130,1510,770,1650),'保存分类习惯\n供下次分类参考',36)
    box(d,(850,1510,1480,1650),'不保存\n本次文件照常保留',36)
    txt(d,130,1770,'打不开或出现错误？',43,True)
    txt(d,130,1860,'保留提示 → 截图给维护人员',42,color=GRAY)
    return im

def clean_style(style,size,bold=False):
    style.font.name='Microsoft YaHei';style.font.size=Pt(size);style.font.bold=bold;style.font.italic=False;style.font.color.rgb=RGBColor(0,0,0)
    rp=style.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    for attr in list(rf.attrib):del rf.attrib[attr]
    for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Microsoft YaHei')
    pp=style.element.find(qn('w:pPr'))
    if pp is not None:
        for child in list(pp):
            if child.tag in [qn('w:pBdr'),qn('w:shd')]:pp.remove(child)

def download_link(doc):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1)
    rid=p.part.relate_to('https://github.com/yang123-pot/wps-rfq-addin/releases/download/v0.1.0/WpsRfqAddin-Setup-0.1.0.exe','http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rid)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'18');rp.append(size)
    color=OxmlElement('w:color');color.set(qn('w:val'),'22744F');rp.append(color);r.append(rp)
    t=OxmlElement('w:t');t.text='点击下载一键安装包';r.append(t);link.append(r);p._p.append(link)
    p.add_run('  需要下载页权限，也可使用维护人员交付的同版EXE。').font.size=Pt(9)

def build():
    FIGS.mkdir(parents=True,exist_ok=True)
    pages=[page1(),page2(),page3(),page4(),page5(),page6(),page7(),page8()]
    titles=['01 安装 只做一次','02 打开材料整理助手','03 选择清单','04 核对材料分类','05 填写品牌和供应商','06 填抬头生成询价总表','07 拆分询价文件','08 保存进度下次继续']
    doc=Document();s=doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=s.right_margin=Cm(2.0);s.top_margin=Cm(1.3);s.bottom_margin=Cm(1.25);s.footer_distance=Cm(.5)
    for st in doc.styles:
        if st.type==1:
            pp=st.element.find(qn('w:pPr'))
            if pp is not None:
                for child in list(pp):
                    if child.tag==qn('w:pBdr'):pp.remove(child)
    clean_style(doc.styles['Normal'],10.5);clean_style(doc.styles['Title'],19,True);clean_style(doc.styles['Heading 1'],19,True)
    doc.styles['Normal'].paragraph_format.space_after=Pt(4)
    doc.styles['Normal'].paragraph_format.line_spacing=1.05
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('材料整理助手 0.1.0   ').font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for i,(im,title) in enumerate(zip(pages,titles),1):
        if i>1:doc.add_page_break()
        h=doc.add_paragraph(title,'Title' if i==1 else 'Heading 1');h.paragraph_format.space_before=Pt(0);h.paragraph_format.space_after=Pt(6)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(2)
        pic=p.add_run().add_picture(str(save(i,im)),width=Cm(16.6))
        pic._inline.docPr.set('descr',title+' 操作示意和界面截图')
        if i==1:
            download_link(doc)
            doc.add_paragraph('适用64位Windows。安装界面为示意；遇到系统安全提示或版本未验证提示，先联系维护人员。').runs[0].font.size=Pt(8)
    doc.core_properties.title='材料整理助手操作图解'
    doc.core_properties.subject='0.1.0 EXE一键安装与七步操作图解 2026年10月21日'
    doc.core_properties.author='材料整理助手'
    doc.save(OUT)
    print(json.dumps({'output':str(OUT),'pages':8,'bytes':OUT.stat().st_size},ensure_ascii=False))

if __name__=='__main__':build()
